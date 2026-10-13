#!/usr/bin/env python3
"""Write the shop page from the spreadsheet.

Prices live in shop-prices.xlsx because that is a thing a person can edit on a
phone between loads of gear, and HTML is not. This reads that sheet, writes
data/shop.json as the machine-readable copy, and rewrites three marked regions
of shop.html: the product cards, the enquiry form's product list, and the
Product nodes in the JSON-LD.

The cards stay *in* the HTML rather than being fetched at runtime. A shop whose
prices arrive by JavaScript is a shop Google prices at nothing, and the whole
point of the structured data is that the price in the rich result matches the
price on the page.

    python3 tools/build_shop.py [--check]

--check exits non-zero if the page is out of date, for CI.
"""

import json
import re
import sys
from datetime import date
from html import escape
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "shop-prices.xlsx"
JSON = ROOT / "data" / "shop.json"
PAGE = ROOT / "shop.html"
SITEMAP = ROOT / "sitemap.xml"

SITE = "https://www.mikidrummer.ca"
SHEET = "Products"

# Column header -> key. Headers are what the spreadsheet shows a human, so they
# are allowed to be wordy; the keys are not.
COLUMNS = {
    "ID":                 "slug",
    "Product name":       "name",
    "Description":        "blurb",
    "Spec headline":      "spec_head",
    "Spec note":          "spec_note",
    "Price":              "price",
    "Currency":           "currency",
    "In stock?":          "in_stock",
    "Checkout link":      "checkout",
    "Image file":         "image",
    "Image description":  "alt",
    "Is the image a rendering?": "rendering",
}


def truthy(v):
    return str(v).strip().lower() in ("yes", "y", "true", "1", "t")


def read_sheet():
    from openpyxl import load_workbook

    wb = load_workbook(XLSX, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit("shop-prices.xlsx has no '%s' tab" % SHEET)
    ws = wb[SHEET]

    rows = list(ws.iter_rows(values_only=True))
    head = [str(c).strip() if c is not None else "" for c in rows[0]]
    missing = [h for h in COLUMNS if h not in head]
    if missing:
        sys.exit("shop-prices.xlsx is missing column(s): " + ", ".join(missing))
    idx = {COLUMNS[h]: head.index(h) for h in COLUMNS}

    out = []
    for r in rows[1:]:
        if not r or not r[idx["slug"]]:
            continue
        get = lambda k: r[idx[k]]
        slug = str(get("slug")).strip()
        if slug.startswith("#"):          # a commented-out row
            continue
        price = get("price")
        if price is None:
            sys.exit("row '%s' has no price" % slug)
        out.append({
            "slug":      slug,
            "name":      str(get("name") or "").strip(),
            "blurb":     str(get("blurb") or "").strip(),
            "spec_head": str(get("spec_head") or "").strip(),
            "spec_note": str(get("spec_note") or "").strip(),
            "price":     round(float(price), 2),
            "currency":  str(get("currency") or "CAD").strip().upper(),
            "in_stock":  truthy(get("in_stock")),
            "checkout":  str(get("checkout") or "").strip(),
            "image":     str(get("image") or "").strip(),
            "alt":       str(get("alt") or "").strip(),
            "rendering": truthy(get("rendering")),
        })
    if not out:
        sys.exit("no product rows found")
    return out


def money(p):
    # $12 reads better than $12.00, but anything with cents keeps both of them:
    # a price shown as $25.5 looks like a typo, which is not what you want on
    # the one number a customer actually reads.
    return "%d" % p if p % 1 == 0 else "%.2f" % p


def card(p):
    e = lambda k: escape(p[k], quote=True)
    gone = "" if p["in_stock"] else " gone"
    bits = ['<article class="item%s" id="%s">' % (gone, e("slug"))]

    if not p["in_stock"]:
        bits.append('      <span class="flag">Sold out</span>')
    bits.append(
        '      <img class="shot" src="/images/shop/%s" width="900" height="630" loading="lazy"\n'
        '           alt="%s">' % (e("image"), e("alt")))
    if p["rendering"]:
        bits.append('      <span class="stand-in">Rendering</span>')

    bits.append('      <div class="body">')
    bits.append('        <div class="name">%s</div>' % e("name"))
    bits.append('        <p class="blurb">%s</p>' % e("blurb"))
    bits.append('        <div class="spec"><b>%s</b>%s</div>' % (e("spec_head"), e("spec_note")))
    bits.append('        <div class="buyrow">')
    bits.append('          <span class="price">$%s <small>%s</small></span>'
                % (money(p["price"]), e("currency")))

    if not p["in_stock"]:
        bits.append('          <button class="buy" type="button" disabled>Sold out</button>')
    elif p["checkout"]:
        # A real checkout when there is a payment link, and never a fake one:
        # with no link the button falls back to the enquiry form rather than
        # pretending at a cart that cannot take money.
        bits.append('          <a class="buy" href="%s" rel="noopener">Buy now</a>' % e("checkout"))
    else:
        bits.append('          <button class="buy" type="button" data-order="%s">Order</button>' % e("name"))
    bits.append('        </div>')

    if not p["in_stock"]:
        bits.append('        <p class="notify">Next batch has no date yet &mdash;\n'
                    '          <button type="button" data-order="%s (restock)">tell me when it is back</button>.</p>'
                    % e("name"))
    bits.append('      </div>')
    bits.append('    </article>')
    return "\n    ".join(bits[:1]) + "\n" + "\n".join(bits[1:])


def product_node(p):
    node = {
        "@type": "Product",
        "name": p["name"],
        "description": p["blurb"],
        "sku": p["slug"],
        "brand": {"@type": "Brand", "name": "Miki Drummer"},
        "url": "%s/shop.html#%s" % (SITE, p["slug"]),
        "image": "%s/images/shop/%s" % (SITE, p["image"]),
        "offers": {
            "@type": "Offer",
            "price": "%.2f" % p["price"],
            "priceCurrency": p["currency"],
            "availability": "https://schema.org/%s"
                            % ("InStock" if p["in_stock"] else "OutOfStock"),
            "url": p["checkout"] or ("%s/shop.html#%s" % (SITE, p["slug"])),
            "seller": {"@type": "ProfessionalService", "name": "Miki Drummer", "url": SITE + "/"},
        },
    }
    if p["spec_head"]:
        node["material"] = p["spec_head"].split("·")[0].strip()
    return node


def between(text, start, end, new):
    pat = re.compile(re.escape(start) + r".*?" + re.escape(end), re.S)
    if not pat.search(text):
        sys.exit("marker pair not found: %s" % start)
    return pat.sub(lambda _: start + "\n" + new + "\n    " + end, text, count=1)


def build(products):
    page = PAGE.read_text(encoding="utf-8")

    cards = "\n\n    ".join(card(p) for p in products)
    page = between(page, "<!-- SHOP:CARDS:START -->", "<!-- SHOP:CARDS:END -->", "    " + cards)

    opts = "\n".join(
        '        <option>%s</option>' % escape(p["name"] + ("" if p["in_stock"] else " (restock)"))
        for p in products)
    page = between(page, "<!-- SHOP:OPTIONS:START -->", "<!-- SHOP:OPTIONS:END -->", opts)

    # The Product nodes sit at the head of the JSON-LD @graph, ahead of the
    # breadcrumb that is written by hand. Parse the block and put the list back
    # rather than splicing text into it: a regex that rewrites JSON has to get
    # the commas right on every run, and the first version of this did not --
    # it left `},,` the second time it ran, which is invalid and silent, since
    # nothing on the page breaks when its structured data stops parsing.
    pat = re.compile(r'(<script type="application/ld\+json">\n)(.*?)(\n</script>)', re.S)
    m = pat.search(page)
    if not m:
        sys.exit("no JSON-LD block in shop.html")
    try:
        doc = json.loads(m.group(2))
    except json.JSONDecodeError as e:
        sys.exit("shop.html's JSON-LD does not parse, fix it by hand first: %s" % e)

    kept = [n for n in doc.get("@graph", []) if n.get("@type") != "Product"]
    doc["@graph"] = [product_node(p) for p in products] + kept
    blob = json.dumps(doc, indent=2, ensure_ascii=False)
    return page[:m.start()] + m.group(1) + blob + m.group(3) + page[m.end():]


def sitemap(out_of_date_ok=True):
    if not SITEMAP.exists():
        return None
    xml = SITEMAP.read_text(encoding="utf-8")
    loc = "%s/shop.html" % SITE
    entry = ("  <url>\n    <loc>%s</loc>\n    <lastmod>%s</lastmod>\n"
             "    <priority>0.8</priority>\n    <changefreq>weekly</changefreq>\n  </url>\n"
             % (loc, date.today().isoformat()))
    if loc in xml:
        return re.sub(r"(<url>\s*<loc>%s</loc>\s*<lastmod>)[^<]*(</lastmod>)" % re.escape(loc),
                      r"\g<1>%s\g<2>" % date.today().isoformat(), xml, count=1)
    return xml.replace("</urlset>", entry + "</urlset>")


def main():
    check = "--check" in sys.argv
    products = read_sheet()

    blob = json.dumps({"updated": date.today().isoformat(), "products": products},
                      indent=2, ensure_ascii=False) + "\n"
    page = build(products)
    xml = sitemap()

    stale = []
    if not JSON.exists() or JSON.read_text(encoding="utf-8") != blob:
        stale.append(JSON.name)
    if PAGE.read_text(encoding="utf-8") != page:
        stale.append(PAGE.name)
    if xml and SITEMAP.read_text(encoding="utf-8") != xml:
        stale.append(SITEMAP.name)

    if check:
        print("stale: " + ", ".join(stale) if stale else "shop is up to date")
        sys.exit(1 if stale else 0)

    JSON.parent.mkdir(exist_ok=True)
    JSON.write_text(blob, encoding="utf-8")
    PAGE.write_text(page, encoding="utf-8")
    if xml:
        SITEMAP.write_text(xml, encoding="utf-8")

    for p in products:
        where = "buy now" if (p["in_stock"] and p["checkout"]) else \
                "enquiry" if p["in_stock"] else "sold out"
        print("  %-14s $%-7s %-9s %s" % (p["slug"], money(p["price"]), where,
                                         "" if (ROOT / "images" / "shop" / p["image"]).exists()
                                         else "!! missing image " + p["image"]))
    print("wrote %s, %s%s" % (JSON.name, PAGE.name, ", " + SITEMAP.name if xml else ""))


if __name__ == "__main__":
    main()
